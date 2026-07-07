from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_PATH = BASE_DIR / "Smart_Mandi_8_Project_Report_Demo.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    if heading.runs:
        heading.runs[0].font.name = "Cambria"
        heading.runs[0].font.size = Pt(16 if level == 1 else 13)
    return heading


def add_paragraph(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        run = paragraph.add_run(bold_prefix)
        run.bold = True
    paragraph.add_run(text)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_title_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PROJECT REPORT")
    r.bold = True
    r.font.size = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ON")
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SMART MANDI 8")
    r.bold = True
    r.font.size = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("(A Django-based Agricultural Market Discovery and Price Information System)")
    r.italic = True
    r.font.size = Pt(13)

    for _ in range(4):
        doc.add_paragraph()

    details = [
        "Submitted in partial fulfillment of the requirements for Semester 5 Project Work",
        "Bachelor of Computer Applications / Bachelor of Science in Information Technology",
        "Academic Year: 2026",
    ]
    for line in details:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line).font.size = Pt(12)

    for _ in range(3):
        doc.add_paragraph()

    info_lines = [
        "Prepared By: ______________________________",
        "Enrollment No.: ___________________________",
        "Guided By: ________________________________",
        "College Name: _____________________________",
    ]
    for line in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(12)

    doc.add_page_break()


def build_preliminary_pages(doc):
    add_heading(doc, "Certificate", level=1)
    add_paragraph(
        doc,
        " This is to certify that the project titled 'Smart Mandi 8' is a bonafide work carried out "
        "for academic demonstration purposes. This report template has been prepared based on the "
        "current project implementation available in the source code repository.",
    )
    doc.add_paragraph()

    add_heading(doc, "Declaration", level=1)
    add_paragraph(
        doc,
        " I hereby declare that the work presented in this report is based on the Smart Mandi 8 project "
        "and has been prepared for demonstration and documentation purposes.",
    )
    doc.add_paragraph()

    add_heading(doc, "Acknowledgement", level=1)
    add_paragraph(
        doc,
        " I express my sincere gratitude to my project guide, department faculty, and peers for their support "
        "and guidance during the development of the Smart Mandi 8 project. Their suggestions helped improve "
        "the design, usability, and structure of this application.",
    )
    doc.add_page_break()


def build_index(doc):
    add_heading(doc, "Index", level=1)
    items = [
        "1. Abstract",
        "2. Introduction",
        "3. Problem Statement",
        "4. Objectives",
        "5. Scope of the Project",
        "6. Technology Stack",
        "7. System Study and Module Description",
        "8. Database Design",
        "9. System Workflow",
        "10. Implementation Details",
        "11. Testing",
        "12. Limitations",
        "13. Future Enhancements",
        "14. Conclusion",
        "15. Bibliography",
    ]
    add_bullets(doc, items)
    doc.add_page_break()


def build_main_content(doc):
    add_heading(doc, "1. Abstract", level=1)
    add_paragraph(
        doc,
        " Smart Mandi 8 is a web-based agricultural market information system developed using Django. "
        "The project helps users search for nearby mandis in Gujarat, view cached and live mandi price data, "
        "register on the platform, submit feedback, send queries, and allow administrators to monitor platform activity. "
        "The system integrates government market-price data, map-based discovery, user account management, and admin analytics "
        "in a single application. The main aim of the project is to present agricultural market information in a cleaner, more "
        "organized, and user-friendly way.",
    )

    add_heading(doc, "2. Introduction", level=1)
    add_paragraph(
        doc,
        " Farmers, traders, and local buyers often need quick access to reliable market and price information. "
        "Traditional methods of collecting mandi data can be slow or inconsistent. Smart Mandi 8 addresses this issue by "
        "combining location-based mandi discovery with current and cached market price records. The application is designed "
        "as a semester project with practical modules such as authentication, feedback collection, query submission, analytics, "
        "and automated data fetching.",
    )

    add_heading(doc, "3. Problem Statement", level=1)
    add_paragraph(
        doc,
        " Agricultural users may face difficulty in identifying nearby mandis and checking available commodity prices in one place. "
        "When market data is scattered across different sources, decision-making becomes slow. A simple digital platform is needed "
        "to centralize mandi discovery, price browsing, and communication features for users.",
    )

    add_heading(doc, "4. Objectives", level=1)
    add_bullets(
        doc,
        [
            "To develop a Django web application for mandi discovery and crop price browsing.",
            "To allow users to search nearby mandis using city input or GPS coordinates.",
            "To display live and cached commodity price information in an organized table.",
            "To provide user registration, login, logout, and profile management.",
            "To store user feedback and support queries in the database.",
            "To provide an admin dashboard for usage and interaction analytics.",
            "To automate daily data fetching and maintain a rolling CSV cache.",
        ],
    )

    add_heading(doc, "5. Scope of the Project", level=1)
    add_paragraph(
        doc,
        " The current project focuses on Gujarat mandi data and web-based access. It provides academic demonstration of "
        "location search, data caching, dashboard reporting, and user interaction features. The project is suitable for "
        "learning Django project structure, CRUD workflows, templates, forms, data integration, and scheduling concepts.",
    )

    add_heading(doc, "6. Technology Stack", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Layer", "Technology", "Purpose"]
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True)
        shade_cell(table.rows[0].cells[idx], "D9EAD3")
    rows = [
        ("Backend", "Python, Django", "Business logic, routing, forms, authentication, ORM"),
        ("Frontend", "HTML, CSS, JavaScript", "User interface and browser interaction"),
        ("Database", "SQLite", "Stores users, profiles, feedback, queries, analytics logs"),
        ("Map & Geo", "Leaflet, OpenStreetMap Nominatim", "Map rendering and city geocoding"),
        ("Scheduler", "APScheduler, django-apscheduler", "Daily automatic mandi data fetch"),
        ("External Data", "data.gov.in API", "Government mandi price records"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)

    add_heading(doc, "7. System Study and Module Description", level=1)
    modules = [
        (
            "Market Finder Module",
            "Provides the landing page, city-based mandi search, GPS-based mandi search, nearby mandi listing, mandi detail view, "
            "commodity filtering, cached snapshots, live API mode, and a debug page for CSV statistics.",
        ),
        (
            "User Module",
            "Handles user registration, login, logout, and profile update functions using Django authentication and a custom profile model.",
        ),
        (
            "About Module",
            "Displays project mission, vision, and key highlights through editable about content.",
        ),
        (
            "Feedback Module",
            "Collects user ratings and feedback messages and stores them in the database.",
        ),
        (
            "Query Module",
            "Allows users to submit project-related queries and keeps track of their status.",
        ),
        (
            "Admin Dashboard Module",
            "Shows user counts, feedback, queries, page visits, mandi search logs, top search cities, and status summaries.",
        ),
        (
            "Scheduler and Data Cache Module",
            "Fetches Gujarat mandi records from the external API, stores them in a rolling CSV file, and purges older rows beyond seven days.",
        ),
    ]
    for title, desc in modules:
        add_paragraph(doc, desc, bold_prefix=f"{title}:")

    add_heading(doc, "8. Database Design", level=1)
    add_paragraph(
        doc,
        " The project uses SQLite as its local database. The main entities are summarized below.",
    )
    db_table = doc.add_table(rows=1, cols=3)
    db_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    db_table.style = "Table Grid"
    db_headers = ["Model", "Important Fields", "Purpose"]
    for idx, header in enumerate(db_headers):
        set_cell_text(db_table.rows[0].cells[idx], header, bold=True)
        shade_cell(db_table.rows[0].cells[idx], "FCE5CD")
    db_rows = [
        ("MandiLocation", "market_name, district, state, latitude, longitude", "Stores mandi coordinates and identity"),
        ("UserProfile", "user, phone_number, address, city, state, pincode", "Extends Django user details"),
        ("Feedback", "name, email, rating, message, created_at", "Stores user feedback"),
        ("UserQuery", "name, email, phone_number, subject, message, status", "Stores support and help requests"),
        ("AboutContent", "title, hero_text, mission, vision, highlights", "Stores About page content"),
        ("FindMandiPageVisit", "user, session_key, ip_address, visited_at", "Tracks finder page visits"),
        ("MandiSearchLog", "city, mandi_name, search_source, latitude, longitude, result_count", "Tracks mandi search activity"),
    ]
    for row in db_rows:
        cells = db_table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)

    add_heading(doc, "9. System Workflow", level=1)
    add_bullets(
        doc,
        [
            "User opens the Smart Mandi home page.",
            "User navigates to the Find Mandi page.",
            "System accepts city input or current GPS location.",
            "Nearby mandis are calculated using the Haversine distance formula.",
            "User opens a mandi detail page to view commodity prices.",
            "System can show cached snapshots or fetch live data from the API.",
            "Registered users can manage profiles and submit feedback or queries.",
            "Admin can review analytics, feedback, and user queries from the dashboard.",
            "The scheduler fetches fresh mandi data every day and updates the rolling CSV cache.",
        ],
    )

    add_heading(doc, "10. Implementation Details", level=1)
    add_paragraph(
        doc,
        " The project is implemented as a multi-app Django system. URL routing is centralized in the main project, while each app "
        "contains its own models, views, forms, templates, and admin configuration. The mandi finder feature uses a 20-kilometer radius "
        "to filter nearby markets from stored latitude and longitude data. Price details support both cached and live modes, which helps "
        "the user compare local snapshots against current API responses. The application also logs page visits and mandi searches for "
        "basic analytical reporting in the admin dashboard.",
    )
    add_paragraph(
        doc,
        " A rolling CSV cache named 'mandi_cache.csv' is maintained beside the Django project root. Daily fetch logic normalizes fields "
        "such as state, district, market, commodity, arrival date, and min/max/modal prices. Old cache rows are purged automatically to "
        "keep the dataset limited to recent snapshots.",
    )

    add_heading(doc, "11. Testing", level=1)
    add_bullets(
        doc,
        [
            "Checked URL routing for home, finder, detail, feedback, query, about, and user pages.",
            "Verified form-based workflows for registration, login, feedback submission, and query submission.",
            "Verified admin dashboard access for staff users through existing test coverage.",
            "Reviewed scheduler integration and manual fetch endpoint for data updates.",
            "Validated template-level handling for empty states, filters, and status messages.",
        ],
    )

    add_heading(doc, "12. Limitations", level=1)
    add_bullets(
        doc,
        [
            "The current implementation is focused on Gujarat mandi data only.",
            "The application depends on third-party APIs for geocoding and mandi price data.",
            "The project uses SQLite, which is suitable for local development but not ideal for large-scale deployment.",
            "Some features are oriented toward academic demonstration rather than production-grade deployment.",
        ],
    )

    add_heading(doc, "13. Future Enhancements", level=1)
    add_bullets(
        doc,
        [
            "Add support for more states and district-wise filters.",
            "Provide charts for commodity price trends over time.",
            "Send notifications when selected commodity prices change significantly.",
            "Add role-based dashboards for farmers, traders, and administrators.",
            "Deploy the project with PostgreSQL and production security settings.",
            "Integrate export options for PDF, Excel, and printable reports.",
        ],
    )

    add_heading(doc, "14. Conclusion", level=1)
    add_paragraph(
        doc,
        " Smart Mandi 8 successfully demonstrates how a Django project can combine agricultural market discovery, live and cached "
        "price viewing, user interaction, and administrative analytics in one platform. The project is practical for academic study "
        "because it covers multiple important web-development concepts such as authentication, forms, database modeling, API integration, "
        "scheduling, and responsive interface design. With additional scaling and production-focused improvements, the project can be "
        "extended into a more complete agri-information portal.",
    )

    add_heading(doc, "15. Bibliography", level=1)
    add_bullets(
        doc,
        [
            "Django official documentation",
            "Python official documentation",
            "data.gov.in agricultural market price API",
            "Leaflet documentation",
            "OpenStreetMap Nominatim documentation",
            "SQLite documentation",
        ],
    )


def apply_document_style(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)


def main():
    doc = Document()
    apply_document_style(doc)
    build_title_page(doc)
    build_preliminary_pages(doc)
    build_index(doc)
    build_main_content(doc)

    section = doc.sections[-1]
    section.start_type = WD_SECTION.NEW_PAGE
    doc.save(OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
